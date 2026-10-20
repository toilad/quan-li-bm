"""The batch result as one .xlsx (owner, 2026-10-20). Sheet KET_QUA keeps the owner's column layout — A name, B/C
empty, D phone, E short name, F FULL ADDRESS, G CITY, H STATE, I ZIP, J file number — with the review columns after
it; sheet CAN_DUYET lists only the REVIEW_REQUIRED records with their reason. Every identifier is TEXT (no 2.01E+11,
no lost leading zero). Written again after each PDF, so a stopped run keeps every row read so far."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fields import short_name

# (header, row key, width, wrap, text)
KET_QUA = [
    ("A · Tên doanh nghiệp", "entity_name", 38, True, True),
    ("B", "", 4, False, True),
    ("C", "", 4, False, True),
    ("D · Điện thoại", "", 14, False, True),
    ("E · Tên rút gọn", "short_name", 26, False, True),
    ("F · FULL ADDRESS", "business_full_address", 46, True, True),
    ("G · CITY", "business_city", 16, False, True),
    ("H · STATE", "business_state", 9, False, True),
    ("I · ZIP", "business_zip", 11, False, True),
    ("J · Mã số", "state_file_number", 16, False, True),
    ("source_file", "source_file", 24, False, True),
    ("check", "check", 26, True, True),
    ("status", "status", 26, False, True),
    ("document_type", "document_type", 18, False, True),
    ("confidence", "confidence", 11, False, False),
    ("reason", "reason", 60, True, True),
    ("address_source", "address_source", 40, True, True),
]
CAN_DUYET = [
    ("source_file", "source_file", 24, False, True),
    ("Tên ứng viên (A)", "entity_name", 38, True, True),
    ("FULL ADDRESS (F)", "business_full_address", 46, True, True),
    ("CITY (G)", "business_city", 16, False, True),
    ("STATE (H)", "business_state", 9, False, True),
    ("ZIP (I)", "business_zip", 11, False, True),
    ("Mã số (J)", "state_file_number", 16, False, True),
    ("reason", "reason", 70, True, True),
    ("document_type", "document_type", 18, False, True),
]
FILL = {"OK": "E8F5E9", "NO_BUSINESS_ADDRESS_ON_FORM": "E3F2FD", "REVIEW_REQUIRED": "FFF4CC", "LOW_CONFIDENCE": "FFF4CC",
        "UNKNOWN_DOCUMENT": "FDE2E1", "OCR_FAILED": "FDE2E1", "ERROR": "FDE2E1"}


def _sheet(ws, columns, rows):
    ws.append([h for h, *_ in columns])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        values = []
        for _, key, _, _, text in columns:
            v = row.get(key, "") if key else ""
            values.append(("" if v is None else str(v)) if text else (v if v not in (None, "") else None))
        ws.append(values)
        r = ws.max_row
        fill = FILL.get(row.get("status", ""))
        for i, (_, _, _, wrap, text) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=i)
            if text:
                cell.number_format = "@"  # TEXT: 200736000006 stays 200736000006, ZIP 02108 keeps its 0
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    for i, (_, _, width, _, text) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if text:  # the empty cells under the data stay TEXT too, for what the owner types in later
            ws.column_dimensions[get_column_letter(i)].number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 1)}"


def write(rows: list[dict], out: Path) -> None:
    for r in rows:
        r.setdefault("short_name", short_name(r.get("entity_name") or "") if r.get("entity_name") else "")
    wb = Workbook()
    ws = wb.active
    ws.title = "KET_QUA"
    _sheet(ws, KET_QUA, rows)
    _sheet(wb.create_sheet("CAN_DUYET"), CAN_DUYET, [r for r in rows if r.get("status") == "REVIEW_REQUIRED"])
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(out)  # never a half-written result file

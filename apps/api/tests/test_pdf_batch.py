"""Folder → one results .xlsx, one row added per PDF (made-up values, no real PDF, no OCR run)."""
from openpyxl import load_workbook

from wg_core.modules.pdf import batch, ocr


def test_one_row_per_pdf_appended_duplicates_skipped_review_keeps_candidates(tmp_path, monkeypatch):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "a.pdf").write_bytes(b"%PDF a")
    (folder / "a - Copy.pdf").write_bytes(b"%PDF a")  # same content, another name
    (folder / "b.pdf").write_bytes(b"%PDF b")
    rows = {b"%PDF a": {"entity_name": "Sample Holdings, LLC", "business_full_address": "1 Sample St, Los Angeles, CA 90001",
                        "business_street": "1 Sample St", "business_city": "Los Angeles", "business_state": "CA",
                        "business_zip": "90001", "state_file_number": "201099900001", "document_type": "LLC-5 REV 04/2007",
                        "status": "OK", "confidence": 97, "reason": "Địa chỉ lấy từ ô mục 7"},
            b"%PDF b": {"entity_name": "Sample Holdings Ill, LLC", "status": "REVIEW_REQUIRED", "reason": "III/Ill"}}
    monkeypatch.setattr(batch, "_text_layer_row", lambda data: None)
    monkeypatch.setattr(batch, "_scan_row", lambda data, force=False: rows[data])
    out = tmp_path / "ket-qua.xlsx"
    s = batch.run(folder, out)
    assert s["files"] == 2 and s["duplicates"] == 1
    wb = load_workbook(out)
    assert wb.sheetnames == ["KET_QUA", "CAN_DUYET"]
    ws = wb["KET_QUA"]
    head = [c.value for c in ws[1]]
    assert head[:10] == ["A · Tên doanh nghiệp", "B", "C", "D · Điện thoại", "E · Tên rút gọn", "F · FULL ADDRESS", "G · CITY",
                         "H · STATE", "I · ZIP", "J · Mã số"]  # the owner's layout: F full, G city, H state, I zip, J number
    a = [c.value for c in ws[2]]
    assert a[0] == "Sample Holdings, LLC" and a[4] == "samplehold" and a[1] == a[2] is None
    assert (a[5], a[6], a[7], a[8]) == ("1 Sample St, Los Angeles, CA 90001", "Los Angeles", "CA", "90001")
    assert a[9] == "201099900001" and isinstance(a[9], str) and ws["J2"].number_format == "@" and ws["I2"].number_format == "@"
    assert a[10] == "a - Copy.pdf" and a[11] == "OK"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["F2"].alignment.wrap_text
    assert not any(ws.column_dimensions[c].hidden for c in "ABCDEFGHIJ")  # the owner's columns shown
    assert all(ws.column_dimensions[c].hidden for c in "KLMNOPQ")  # review columns after J start hidden
    b = [c.value for c in ws[3]]
    assert b[0] == "Sample Holdings Ill, LLC" and b[12] == "REVIEW_REQUIRED" and b[11].startswith("CẦN DUYỆT")  # candidate kept
    review = [[c.value for c in r] for r in wb["CAN_DUYET"].iter_rows(min_row=2)]
    rh = [c.value for c in wb["CAN_DUYET"][1]]
    assert {"check", "status", "document_type", "confidence", "reason", "address_source"} <= set(rh)  # every debug column
    assert [(r[0], r[1], r[rh.index("reason")]) for r in review] == [("b.pdf", "Sample Holdings Ill, LLC", "III/Ill")]  # review only


def test_ocr_cache_is_small_json_capped_and_clearable(tmp_path, monkeypatch):
    monkeypatch.setenv("OCR_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setenv("OCR_CACHE_LIMIT_MB", "10")
    d = tmp_path / "cache" / "ab"
    d.mkdir(parents=True)
    (d / "ab-p1-original-x.json").write_text("{}")
    (d / "ab-p1-threshold-x.json").write_text("{}")  # an old whole-page variant: no longer read
    big = d / "ab-p2-original-x.json"
    big.write_bytes(b" " * 11_000_000)
    import os, time  # noqa: E401

    os.utime(big, (time.time() - 100, time.time() - 100))  # the oldest page goes first
    assert ocr.prune_cache() == 2
    assert [f.name for f in d.iterdir()] == ["ab-p1-original-x.json"]
    assert ocr.cache_info()["mode"] == "minimal" and ocr.clear_cache() == 1 and ocr.cache_info()["files"] == 0


def test_files_and_folders_picked_together_unknown_gets_no_guess(tmp_path, monkeypatch):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "x.pdf").write_bytes(b"%PDF x")
    (folder / "notes.txt").write_text("not a pdf")
    single = tmp_path / "y.PDF"
    single.write_bytes(b"%PDF y")
    rows = {b"%PDF x": {"status": "UNKNOWN_DOCUMENT", "entity_name": "guess", "reason": "Không khớp mẫu"},
            b"%PDF y": {"status": "NO_BUSINESS_ADDRESS_ON_FORM", "entity_name": "Sample LLC", "state_file_number": "201099900001"}}
    monkeypatch.setattr(batch, "_text_layer_row", lambda data: None)
    monkeypatch.setattr(batch, "_scan_row", lambda data, force=False: rows[data])
    out = tmp_path / "r.xlsx"
    s = batch.run([folder, single], out)
    ws = load_workbook(out)["KET_QUA"]
    got = {r[10].value: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    assert s["files"] == 2 and set(got) == {"x.pdf", "y.PDF"}
    assert got["x.pdf"][0] in (None, "") and got["x.pdf"][11].startswith("KHÔNG ĐỌC")  # nothing guessed
    assert got["y.PDF"][5] in (None, "") and got["y.PDF"][11].startswith("OK") and got["y.PDF"][9] == "201099900001"
